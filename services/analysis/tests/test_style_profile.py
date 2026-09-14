"""build_style_profile (issue #96): deterministic font/color/margin/
one-page-fit extraction from a PDF/DOCX CVVersion's raw bytes, plus one LLM
call classifying the layout archetype and each heading's SectionType (+
sidebar/main region). See docs/adr/0008-tailored-cv-visual-style-profile.md.
"""

import io

import pytest
from docx import Document as DocxDocument
from py_db.models import Cvfiletype

from analysis.llm_provider import LLMProvider
from analysis.style_profile import StyleProfileError, build_style_profile

MARKDOWN = "# Experience\n\n- Senior Backend Engineer, Acme Corp\n\n# Skills\n\n- Python\n- AWS"

CLASSIFICATION_JSON = (
    '{"layoutArchetype": "SINGLE_COLUMN", "sections": '
    '[{"heading": "Experience", "sectionType": "EXPERIENCE", "region": null}, '
    '{"heading": "Skills", "sectionType": "SKILLS", "region": null}]}'
)

SIDEBAR_CLASSIFICATION_JSON = (
    '{"layoutArchetype": "SIDEBAR_MAIN", "sections": '
    '[{"heading": "Experience", "sectionType": "EXPERIENCE", "region": "MAIN"}, '
    '{"heading": "Skills", "sectionType": "SKILLS", "region": "SIDEBAR"}]}'
)


class StubLLMProvider(LLMProvider):
    def __init__(self, response: str):
        self._response = response
        self.calls = 0

    def generate(self, *, system: str, prompt: str, max_tokens: int | None = None) -> str:
        self.calls += 1
        return self._response


def _build_fixture_pdf_bytes() -> bytes:
    """A single-page PDF with a larger red heading line and a smaller black
    body line, so heading vs. body font size/color extraction has something
    real to distinguish."""
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /Resources << /Font << /F1 4 0 R >> >> "
        b"/MediaBox [0 0 612 792] /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    content = (
        b"BT /F1 20 Tf 1 0 0 rg 100 700 Td (Experience) Tj ET\n"
        b"BT /F1 11 Tf 0 0 0 rg 100 650 Td "
        b"(Senior Backend Engineer Acme Corp 2019 2024) Tj ET"
    )
    objects.append(b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream")

    buf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(buf))
        buf += f"{i} 0 obj\n".encode()
        buf += obj
        buf += b"\nendobj\n"
    xref_offset = len(buf)
    buf += f"xref\n0 {len(objects) + 1}\n".encode()
    buf += b"0000000000 65535 f\r\n"
    for off in offsets[1:]:
        buf += f"{off:010d} 00000 n\r\n".encode()
    buf += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF".encode()
    return bytes(buf)


def _build_fixture_docx_bytes() -> bytes:
    document = DocxDocument()
    document.add_paragraph("Experience", style="Heading 1")
    document.add_paragraph("Senior Backend Engineer, Acme Corp (2019-2024)")
    document.add_paragraph("Skills", style="Heading 1")
    document.add_paragraph("Python, AWS, PostgreSQL")
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


FIXTURE_PDF_BYTES = _build_fixture_pdf_bytes()
FIXTURE_DOCX_BYTES = _build_fixture_docx_bytes()


def test_build_style_profile_extracts_distinct_pdf_heading_and_body_style():
    provider = StubLLMProvider(CLASSIFICATION_JSON)

    profile = build_style_profile(
        file_bytes=FIXTURE_PDF_BYTES,
        file_type=Cvfiletype.PDF,
        markdown=MARKDOWN,
        llm_provider=provider,
    )

    assert profile["layoutArchetype"] == "SINGLE_COLUMN"
    assert profile["fonts"]["name"] == "Helvetica"
    assert profile["fonts"]["family"] == "sans-serif"
    assert profile["accentColor"] == "#000000"
    assert profile["onePageFit"] is True
    assert set(profile["margins"]) == {"top", "bottom", "left", "right"}
    assert profile["sections"]["EXPERIENCE"]["headingTreatment"]["color"] == "#FF0000"
    assert profile["sections"]["EXPERIENCE"]["region"] is None
    assert profile["sections"]["SKILLS"]["region"] is None
    assert provider.calls == 1


def test_build_style_profile_extracts_docx_margins_and_fonts():
    provider = StubLLMProvider(SIDEBAR_CLASSIFICATION_JSON)

    profile = build_style_profile(
        file_bytes=FIXTURE_DOCX_BYTES,
        file_type=Cvfiletype.DOCX,
        markdown=MARKDOWN,
        llm_provider=provider,
    )

    assert profile["layoutArchetype"] == "SIDEBAR_MAIN"
    assert profile["margins"] == {"top": 72.0, "bottom": 72.0, "left": 90.0, "right": 90.0}
    assert profile["fonts"]["name"] == "Calibri"
    assert profile["onePageFit"] is True
    assert profile["sections"]["EXPERIENCE"]["region"] == "MAIN"
    assert profile["sections"]["SKILLS"]["region"] == "SIDEBAR"


def test_build_style_profile_raises_on_malformed_llm_classification():
    provider = StubLLMProvider("not json")

    with pytest.raises(StyleProfileError):
        build_style_profile(
            file_bytes=FIXTURE_PDF_BYTES,
            file_type=Cvfiletype.PDF,
            markdown=MARKDOWN,
            llm_provider=provider,
        )


def test_build_style_profile_raises_when_markdown_has_no_headings():
    provider = StubLLMProvider(CLASSIFICATION_JSON)

    with pytest.raises(StyleProfileError):
        build_style_profile(
            file_bytes=FIXTURE_PDF_BYTES,
            file_type=Cvfiletype.PDF,
            markdown="just a paragraph, no headings",
            llm_provider=provider,
        )

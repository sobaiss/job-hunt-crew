import asyncio
import json
import uuid
from datetime import UTC, datetime

from io import BytesIO

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from py_db.models import (
    Analysis,
    Analysisstatus,
    CVVersion,
    Cvstylestatus,
    GeneratedDocument,
    Generateddocumentstatus,
    JobOffer,
    Joboffersourcesite,
    Plan,
    PlanQuotaDefault,
    Quotakind,
    User,
)
from py_db.session import make_engine, make_session_factory
from sqlalchemy import delete, select, update

from api.main import app
from api.sqs_client import GENERATION_INTAKE_QUEUE_URL, make_sqs_client

INTERNAL_SECRET_HEADERS = {"X-Internal-Api-Secret": "test-secret"}


@pytest.fixture(autouse=True)
def _secret(monkeypatch):
    monkeypatch.setenv("INTERNAL_API_SECRET", "test-secret")


def _headers(user_id: str) -> dict[str, str]:
    return {**INTERNAL_SECRET_HEADERS, "X-User-Id": user_id}


def _now() -> datetime:
    return datetime.now(UTC).replace(tzinfo=None)


async def _create_user() -> str:
    user_id = str(uuid.uuid4())
    engine = make_engine()
    try:
        session_factory = make_session_factory(engine)
        async with session_factory() as session:
            session.add(User(id=user_id, email=f"{user_id}@example.com", updatedAt=_now()))
            await session.commit()
    finally:
        await engine.dispose()
    return user_id


async def _delete_user(user_id: str) -> None:
    engine = make_engine()
    try:
        session_factory = make_session_factory(engine)
        async with session_factory() as session:
            # GeneratedDocument.cvVersionId is onDelete: Restrict — clear any
            # rows this test created before the CVVersion cascade from User.
            cv_version_ids = (
                await session.scalars(select(CVVersion.id).where(CVVersion.userId == user_id))
            ).all()
            if cv_version_ids:
                await session.execute(
                    delete(GeneratedDocument).where(
                        GeneratedDocument.cvVersionId.in_(cv_version_ids)
                    )
                )
            await session.execute(delete(User).where(User.id == user_id))
            await session.commit()
    finally:
        await engine.dispose()


@pytest.fixture
def user_id():
    uid = asyncio.run(_create_user())
    yield uid
    asyncio.run(_delete_user(uid))


async def _free_plan_default(kind: Quotakind) -> int | None:
    engine = make_engine()
    try:
        session_factory = make_session_factory(engine)
        async with session_factory() as session:
            return await session.scalar(
                select(PlanQuotaDefault.limit).where(
                    PlanQuotaDefault.plan == Plan.FREE, PlanQuotaDefault.quotaKind == kind
                )
            )
    finally:
        await engine.dispose()


async def _set_free_plan_default(kind: Quotakind, limit: int | None) -> None:
    engine = make_engine()
    try:
        session_factory = make_session_factory(engine)
        async with session_factory() as session:
            await session.execute(
                update(PlanQuotaDefault)
                .where(PlanQuotaDefault.plan == Plan.FREE, PlanQuotaDefault.quotaKind == kind)
                .values(limit=limit)
            )
            await session.commit()
    finally:
        await engine.dispose()


@pytest.fixture
def documents_daily_cap():
    """Temporarily overrides FREE's `PlanQuotaDefault` for `DOCUMENTS_DAILY`
    — `user_id` above creates a plain FREE-plan User (issue #135), so this is
    the fixture cap tests use instead of the retired `DAILY_GENERATION_CAP`
    env var (issue #137), matching test_v1_analyses.py's `analyses_daily_cap`
    pattern.
    """
    original = asyncio.run(_free_plan_default(Quotakind.DOCUMENTS_DAILY))

    def _apply(limit: int | None) -> None:
        asyncio.run(_set_free_plan_default(Quotakind.DOCUMENTS_DAILY, limit))

    yield _apply
    asyncio.run(_set_free_plan_default(Quotakind.DOCUMENTS_DAILY, original))


def _drain_generation_intake() -> list[dict]:
    sqs = make_sqs_client()
    bodies: list[dict] = []
    while True:
        received = sqs.receive_message(QueueUrl=GENERATION_INTAKE_QUEUE_URL, MaxNumberOfMessages=10)
        messages = received.get("Messages", [])
        if not messages:
            break
        for message in messages:
            bodies.append(json.loads(message["Body"]))
            sqs.delete_message(
                QueueUrl=GENERATION_INTAKE_QUEUE_URL, ReceiptHandle=message["ReceiptHandle"]
            )
    return bodies


@pytest.fixture(autouse=True)
def _clean_generation_intake_queue():
    _drain_generation_intake()
    yield
    _drain_generation_intake()


def _make_cv_version(client: TestClient, user_id: str) -> str:
    response = client.post(
        "/v1/cv-versions",
        headers=_headers(user_id),
        json={
            "label": "Base CV",
            "fileName": "cv.pdf",
            "contentType": "application/pdf",
            "fileSizeBytes": 1024,
        },
    )
    assert response.status_code == 201
    return response.json()["cvVersionId"]


async def _seed_analysis(*, user_id: str, cv_version_id: str, status: Analysisstatus) -> str:
    job_offer_id = str(uuid.uuid4())
    analysis_id = str(uuid.uuid4())
    engine = make_engine()
    try:
        session_factory = make_session_factory(engine)
        async with session_factory() as session:
            session.add(
                JobOffer(
                    id=job_offer_id,
                    sourceUrl=f"https://example.com/jobs/{job_offer_id}",
                    sourceSite=Joboffersourcesite.OTHER,
                    updatedAt=_now(),
                )
            )
            session.add(
                Analysis(
                    id=analysis_id,
                    userId=user_id,
                    jobOfferId=job_offer_id,
                    cvVersionId=cv_version_id,
                    status=status,
                    resultJSON={"matched_skills": [], "missing_skills": []}
                    if status == Analysisstatus.COMPLETED
                    else None,
                    matchScore=80 if status == Analysisstatus.COMPLETED else None,
                    completedAt=_now() if status == Analysisstatus.COMPLETED else None,
                )
            )
            await session.commit()
    finally:
        await engine.dispose()
    return analysis_id


def test_create_generated_documents_creates_both_types_and_enqueues(user_id):
    with TestClient(app) as client:
        cv = _make_cv_version(client, user_id)
        analysis_id = asyncio.run(
            _seed_analysis(user_id=user_id, cv_version_id=cv, status=Analysisstatus.COMPLETED)
        )

        response = client.post(
            f"/v1/analyses/{analysis_id}/generated-documents", headers=_headers(user_id)
        )
        assert response.status_code == 202
        body = response.json()
        types = {doc["type"] for doc in body["generatedDocuments"]}
        assert types == {"COVER_LETTER", "TAILORED_CV"}
        for doc in body["generatedDocuments"]:
            assert doc["status"] == "PENDING"
            assert doc["analysisId"] == analysis_id

        enqueued = _drain_generation_intake()
        assert {b["generatedDocumentId"] for b in enqueued} == {
            doc["id"] for doc in body["generatedDocuments"]
        }


def test_create_generated_documents_rejects_a_non_completed_analysis(user_id):
    with TestClient(app) as client:
        cv = _make_cv_version(client, user_id)
        analysis_id = asyncio.run(
            _seed_analysis(user_id=user_id, cv_version_id=cv, status=Analysisstatus.PENDING)
        )

        response = client.post(
            f"/v1/analyses/{analysis_id}/generated-documents", headers=_headers(user_id)
        )
        assert response.status_code == 400
        assert _drain_generation_intake() == []


def test_create_generated_documents_creates_an_application(user_id):
    with TestClient(app) as client:
        cv = _make_cv_version(client, user_id)
        analysis_id = asyncio.run(
            _seed_analysis(user_id=user_id, cv_version_id=cv, status=Analysisstatus.COMPLETED)
        )

        response = client.post(
            f"/v1/analyses/{analysis_id}/generated-documents", headers=_headers(user_id)
        )
        assert response.status_code == 202

        applications = client.get("/v1/applications", headers=_headers(user_id)).json()[
            "applications"
        ]
        matching = [a for a in applications if a["analysisId"] == analysis_id]
        assert len(matching) == 1
        assert matching[0]["status"] == "DRAFT"


def test_create_generated_documents_does_not_duplicate_an_existing_application(user_id):
    with TestClient(app) as client:
        cv = _make_cv_version(client, user_id)
        analysis_id = asyncio.run(
            _seed_analysis(user_id=user_id, cv_version_id=cv, status=Analysisstatus.COMPLETED)
        )

        created = client.post(
            "/v1/applications", headers=_headers(user_id), json={"analysisId": analysis_id}
        ).json()["application"]

        client.post(f"/v1/analyses/{analysis_id}/generated-documents", headers=_headers(user_id))

        applications = client.get("/v1/applications", headers=_headers(user_id)).json()[
            "applications"
        ]
        matching = [a for a in applications if a["analysisId"] == analysis_id]
        assert len(matching) == 1
        assert matching[0]["id"] == created["id"]


def test_create_generated_documents_is_user_scoped(user_id):
    other = asyncio.run(_create_user())
    try:
        with TestClient(app) as client:
            cv = _make_cv_version(client, user_id)
            analysis_id = asyncio.run(
                _seed_analysis(user_id=user_id, cv_version_id=cv, status=Analysisstatus.COMPLETED)
            )
            response = client.post(
                f"/v1/analyses/{analysis_id}/generated-documents", headers=_headers(other)
            )
        assert response.status_code == 404
    finally:
        asyncio.run(_delete_user(other))


def test_list_generated_documents_returns_current_documents_for_the_analysis(user_id):
    with TestClient(app) as client:
        cv = _make_cv_version(client, user_id)
        analysis_id = asyncio.run(
            _seed_analysis(user_id=user_id, cv_version_id=cv, status=Analysisstatus.COMPLETED)
        )
        created = client.post(
            f"/v1/analyses/{analysis_id}/generated-documents", headers=_headers(user_id)
        ).json()

        response = client.get(
            f"/v1/analyses/{analysis_id}/generated-documents", headers=_headers(user_id)
        )
        assert response.status_code == 200
        ids = {doc["id"] for doc in response.json()["generatedDocuments"]}
        assert ids == {doc["id"] for doc in created["generatedDocuments"]}


def test_list_generated_documents_is_empty_before_generation(user_id):
    with TestClient(app) as client:
        cv = _make_cv_version(client, user_id)
        analysis_id = asyncio.run(
            _seed_analysis(user_id=user_id, cv_version_id=cv, status=Analysisstatus.COMPLETED)
        )

        response = client.get(
            f"/v1/analyses/{analysis_id}/generated-documents", headers=_headers(user_id)
        )
        assert response.status_code == 200
        assert response.json()["generatedDocuments"] == []


def test_list_generated_documents_is_user_scoped(user_id):
    other = asyncio.run(_create_user())
    try:
        with TestClient(app) as client:
            cv = _make_cv_version(client, user_id)
            analysis_id = asyncio.run(
                _seed_analysis(user_id=user_id, cv_version_id=cv, status=Analysisstatus.COMPLETED)
            )
            response = client.get(
                f"/v1/analyses/{analysis_id}/generated-documents", headers=_headers(other)
            )
        assert response.status_code == 404
    finally:
        asyncio.run(_delete_user(other))


def test_get_generated_document_returns_the_row(user_id):
    with TestClient(app) as client:
        cv = _make_cv_version(client, user_id)
        analysis_id = asyncio.run(
            _seed_analysis(user_id=user_id, cv_version_id=cv, status=Analysisstatus.COMPLETED)
        )
        created = client.post(
            f"/v1/analyses/{analysis_id}/generated-documents", headers=_headers(user_id)
        ).json()
        document_id = created["generatedDocuments"][0]["id"]

        response = client.get(f"/v1/generated-documents/{document_id}", headers=_headers(user_id))
        assert response.status_code == 200
        assert response.json()["generatedDocument"]["id"] == document_id


def test_get_generated_document_is_user_scoped(user_id):
    other = asyncio.run(_create_user())
    try:
        with TestClient(app) as client:
            cv = _make_cv_version(client, user_id)
            analysis_id = asyncio.run(
                _seed_analysis(user_id=user_id, cv_version_id=cv, status=Analysisstatus.COMPLETED)
            )
            created = client.post(
                f"/v1/analyses/{analysis_id}/generated-documents", headers=_headers(user_id)
            ).json()
            document_id = created["generatedDocuments"][0]["id"]

            response = client.get(
                f"/v1/generated-documents/{document_id}", headers=_headers(other)
            )
        assert response.status_code == 404
    finally:
        asyncio.run(_delete_user(other))


def test_get_generated_document_missing_id_is_404(user_id):
    with TestClient(app) as client:
        response = client.get(
            f"/v1/generated-documents/{uuid.uuid4()}", headers=_headers(user_id)
        )
        assert response.status_code == 404


async def _mark_document_ready(document_id: str, markdown_content: str) -> None:
    engine = make_engine()
    try:
        session_factory = make_session_factory(engine)
        async with session_factory() as session:
            document = await session.get(GeneratedDocument, document_id)
            document.status = Generateddocumentstatus.READY
            document.markdownContent = markdown_content
            await session.commit()
    finally:
        await engine.dispose()


async def _mark_cv_version_style(
    cv_version_id: str, *, style_status: Cvstylestatus, style_profile: dict | None
) -> None:
    engine = make_engine()
    try:
        session_factory = make_session_factory(engine)
        async with session_factory() as session:
            cv_version = await session.get(CVVersion, cv_version_id)
            cv_version.styleStatus = style_status
            cv_version.styleProfile = style_profile
            await session.commit()
    finally:
        await engine.dispose()


_TAILORED_CV_STYLE_PROFILE = {
    "layoutArchetype": "SINGLE_COLUMN",
    "fonts": {"name": "Georgia", "family": "serif"},
    "accentColor": "#112233",
    "margins": {"top": 40, "bottom": 40, "left": 40, "right": 40},
    "onePageFit": True,
    "photoAssetRef": None,
    "sections": {
        "EXPERIENCE": {
            "headingTreatment": {"font": {"name": "Impact", "family": "sans-serif"}, "color": "#FF0000"},
            "region": None,
        }
    },
}


def test_get_generated_document_download_applies_style_profile_to_tailored_cv_docx(user_id):
    with TestClient(app) as client:
        cv = _make_cv_version(client, user_id)
        analysis_id = asyncio.run(
            _seed_analysis(user_id=user_id, cv_version_id=cv, status=Analysisstatus.COMPLETED)
        )
        created = client.post(
            f"/v1/analyses/{analysis_id}/generated-documents", headers=_headers(user_id)
        ).json()
        tailored_cv_id = next(
            doc["id"] for doc in created["generatedDocuments"] if doc["type"] == "TAILORED_CV"
        )
        asyncio.run(
            _mark_document_ready(
                tailored_cv_id, "## Experience\n<!-- SectionType: EXPERIENCE -->\n\nDid things.\n"
            )
        )
        asyncio.run(
            _mark_cv_version_style(
                cv, style_status=Cvstylestatus.EXTRACTED, style_profile=_TAILORED_CV_STYLE_PROFILE
            )
        )

        response = client.get(
            f"/v1/generated-documents/{tailored_cv_id}/download?format=docx",
            headers=_headers(user_id),
        )
        assert response.status_code == 200
        document = DocxDocument(BytesIO(response.content))
        heading_paragraph = next(p for p in document.paragraphs if p.text == "Experience")
        # "Impact" isn't in the curated cross-platform font set, so it's
        # substituted by its family's ("sans-serif") entry instead (#100).
        assert heading_paragraph.runs[0].font.name == "Arial"
        assert str(heading_paragraph.runs[0].font.color.rgb) == "FF0000"


def test_get_generated_document_download_falls_back_to_generic_template_without_extracted_style(
    user_id,
):
    with TestClient(app) as client:
        cv = _make_cv_version(client, user_id)
        analysis_id = asyncio.run(
            _seed_analysis(user_id=user_id, cv_version_id=cv, status=Analysisstatus.COMPLETED)
        )
        created = client.post(
            f"/v1/analyses/{analysis_id}/generated-documents", headers=_headers(user_id)
        ).json()
        tailored_cv_id = next(
            doc["id"] for doc in created["generatedDocuments"] if doc["type"] == "TAILORED_CV"
        )
        asyncio.run(
            _mark_document_ready(
                tailored_cv_id, "## Experience\n<!-- SectionType: EXPERIENCE -->\n\nDid things.\n"
            )
        )
        # cv_version.styleStatus stays PENDING (never extracted) — rendering
        # must fall back to the generic template rather than erroring.

        response = client.get(
            f"/v1/generated-documents/{tailored_cv_id}/download?format=docx",
            headers=_headers(user_id),
        )
        assert response.status_code == 200
        document = DocxDocument(BytesIO(response.content))
        heading_paragraph = next(p for p in document.paragraphs if p.text == "Experience")
        assert heading_paragraph.runs[0].font.name is None


def test_get_generated_document_download_strips_section_type_tag(user_id):
    with TestClient(app) as client:
        cv = _make_cv_version(client, user_id)
        analysis_id = asyncio.run(
            _seed_analysis(user_id=user_id, cv_version_id=cv, status=Analysisstatus.COMPLETED)
        )
        created = client.post(
            f"/v1/analyses/{analysis_id}/generated-documents", headers=_headers(user_id)
        ).json()
        tailored_cv_id = next(
            doc["id"] for doc in created["generatedDocuments"] if doc["type"] == "TAILORED_CV"
        )
        asyncio.run(
            _mark_document_ready(
                tailored_cv_id, "## Experience\n<!-- SectionType: EXPERIENCE -->\n\nDid things.\n"
            )
        )

        response = client.get(
            f"/v1/generated-documents/{tailored_cv_id}/download?format=txt",
            headers=_headers(user_id),
        )
        assert response.status_code == 200
        assert "SectionType" not in response.content.decode("utf-8")


def test_get_generated_document_download_defaults_to_pdf(user_id):
    with TestClient(app) as client:
        cv = _make_cv_version(client, user_id)
        analysis_id = asyncio.run(
            _seed_analysis(user_id=user_id, cv_version_id=cv, status=Analysisstatus.COMPLETED)
        )
        created = client.post(
            f"/v1/analyses/{analysis_id}/generated-documents", headers=_headers(user_id)
        ).json()
        document_id = created["generatedDocuments"][0]["id"]
        asyncio.run(_mark_document_ready(document_id, "# Cover Letter\n\nDear hiring manager."))

        response = client.get(
            f"/v1/generated-documents/{document_id}/download", headers=_headers(user_id)
        )
        assert response.status_code == 200
        assert response.headers["content-type"] == "application/pdf"
        assert response.content.startswith(b"%PDF")


@pytest.mark.parametrize(
    ("format", "content_type", "signature_check"),
    [
        ("pdf", "application/pdf", lambda content: content.startswith(b"%PDF")),
        (
            "docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            lambda content: content.startswith(b"PK"),
        ),
        (
            "md",
            "text/markdown; charset=utf-8",
            lambda content: content.decode("utf-8") == "# Cover Letter\n\nDear hiring manager.",
        ),
        (
            "txt",
            "text/plain; charset=utf-8",
            lambda content: "# Cover Letter" not in content.decode("utf-8"),
        ),
    ],
)
def test_get_generated_document_download_supports_every_format(
    user_id, format, content_type, signature_check
):
    with TestClient(app) as client:
        cv = _make_cv_version(client, user_id)
        analysis_id = asyncio.run(
            _seed_analysis(user_id=user_id, cv_version_id=cv, status=Analysisstatus.COMPLETED)
        )
        created = client.post(
            f"/v1/analyses/{analysis_id}/generated-documents", headers=_headers(user_id)
        ).json()
        document_id = created["generatedDocuments"][0]["id"]
        asyncio.run(_mark_document_ready(document_id, "# Cover Letter\n\nDear hiring manager."))

        response = client.get(
            f"/v1/generated-documents/{document_id}/download?format={format}",
            headers=_headers(user_id),
        )
        assert response.status_code == 200
        assert response.headers["content-type"] == content_type
        assert response.headers["content-disposition"].endswith(f'.{format}"')
        assert signature_check(response.content)


@pytest.mark.parametrize("format", ["pdf", "docx", "md", "txt"])
def test_get_generated_document_download_rejects_a_document_not_yet_ready(user_id, format):
    with TestClient(app) as client:
        cv = _make_cv_version(client, user_id)
        analysis_id = asyncio.run(
            _seed_analysis(user_id=user_id, cv_version_id=cv, status=Analysisstatus.COMPLETED)
        )
        created = client.post(
            f"/v1/analyses/{analysis_id}/generated-documents", headers=_headers(user_id)
        ).json()
        document_id = created["generatedDocuments"][0]["id"]

        response = client.get(
            f"/v1/generated-documents/{document_id}/download?format={format}",
            headers=_headers(user_id),
        )
        assert response.status_code == 400


@pytest.mark.parametrize("format", ["pdf", "docx", "md", "txt"])
def test_get_generated_document_download_is_user_scoped(user_id, format):
    other = asyncio.run(_create_user())
    try:
        with TestClient(app) as client:
            cv = _make_cv_version(client, user_id)
            analysis_id = asyncio.run(
                _seed_analysis(user_id=user_id, cv_version_id=cv, status=Analysisstatus.COMPLETED)
            )
            created = client.post(
                f"/v1/analyses/{analysis_id}/generated-documents", headers=_headers(user_id)
            ).json()
            document_id = created["generatedDocuments"][0]["id"]
            asyncio.run(_mark_document_ready(document_id, "Some content."))

            response = client.get(
                f"/v1/generated-documents/{document_id}/download?format={format}",
                headers=_headers(other),
            )
        assert response.status_code == 404
    finally:
        asyncio.run(_delete_user(other))


def test_regenerate_generated_document_supersedes_and_enqueues(user_id):
    with TestClient(app) as client:
        cv = _make_cv_version(client, user_id)
        analysis_id = asyncio.run(
            _seed_analysis(user_id=user_id, cv_version_id=cv, status=Analysisstatus.COMPLETED)
        )
        created = client.post(
            f"/v1/analyses/{analysis_id}/generated-documents", headers=_headers(user_id)
        ).json()
        document_id = created["generatedDocuments"][0]["id"]
        doc_type = created["generatedDocuments"][0]["type"]
        asyncio.run(_mark_document_ready(document_id, "# Cover Letter\n\nDear hiring manager."))
        _drain_generation_intake()

        response = client.post(
            f"/v1/generated-documents/{document_id}/regenerate", headers=_headers(user_id)
        )
        assert response.status_code == 202
        new_document = response.json()["generatedDocument"]
        assert new_document["id"] != document_id
        assert new_document["type"] == doc_type
        assert new_document["status"] == "PENDING"

        enqueued = _drain_generation_intake()
        assert {b["generatedDocumentId"] for b in enqueued} == {new_document["id"]}

        listed = client.get(
            f"/v1/analyses/{analysis_id}/generated-documents", headers=_headers(user_id)
        ).json()["generatedDocuments"]
        listed_ids = {doc["id"] for doc in listed}
        assert new_document["id"] in listed_ids
        assert document_id not in listed_ids


def test_regenerate_generated_document_rejects_one_already_in_progress(user_id):
    with TestClient(app) as client:
        cv = _make_cv_version(client, user_id)
        analysis_id = asyncio.run(
            _seed_analysis(user_id=user_id, cv_version_id=cv, status=Analysisstatus.COMPLETED)
        )
        created = client.post(
            f"/v1/analyses/{analysis_id}/generated-documents", headers=_headers(user_id)
        ).json()
        document_id = created["generatedDocuments"][0]["id"]
        _drain_generation_intake()

        response = client.post(
            f"/v1/generated-documents/{document_id}/regenerate", headers=_headers(user_id)
        )
        assert response.status_code == 409
        assert _drain_generation_intake() == []


def test_regenerate_generated_document_is_user_scoped(user_id):
    other = asyncio.run(_create_user())
    try:
        with TestClient(app) as client:
            cv = _make_cv_version(client, user_id)
            analysis_id = asyncio.run(
                _seed_analysis(user_id=user_id, cv_version_id=cv, status=Analysisstatus.COMPLETED)
            )
            created = client.post(
                f"/v1/analyses/{analysis_id}/generated-documents", headers=_headers(user_id)
            ).json()
            document_id = created["generatedDocuments"][0]["id"]
            asyncio.run(_mark_document_ready(document_id, "Some content."))

            response = client.post(
                f"/v1/generated-documents/{document_id}/regenerate", headers=_headers(other)
            )
        assert response.status_code == 404
    finally:
        asyncio.run(_delete_user(other))


def test_regenerate_generated_document_respects_daily_generation_cap(
    documents_daily_cap, user_id
):
    documents_daily_cap(2)
    with TestClient(app) as client:
        cv = _make_cv_version(client, user_id)
        analysis_id = asyncio.run(
            _seed_analysis(user_id=user_id, cv_version_id=cv, status=Analysisstatus.COMPLETED)
        )
        created = client.post(
            f"/v1/analyses/{analysis_id}/generated-documents", headers=_headers(user_id)
        ).json()
        # The initial create already produced 2 rows (COVER_LETTER + TAILORED_CV),
        # exhausting a cap of 2 for the day.
        document_id = created["generatedDocuments"][0]["id"]
        asyncio.run(_mark_document_ready(document_id, "Some content."))
        _drain_generation_intake()

        response = client.post(
            f"/v1/generated-documents/{document_id}/regenerate", headers=_headers(user_id)
        )
        assert response.status_code == 429


def test_create_generated_documents_returns_429_once_daily_cap_reached(
    documents_daily_cap, user_id
):
    # A single create call produces 2 rows (COVER_LETTER + TAILORED_CV), so a
    # cap of 2 is already exhausted by the first call.
    documents_daily_cap(2)
    with TestClient(app) as client:
        cv = _make_cv_version(client, user_id)
        first_analysis_id = asyncio.run(
            _seed_analysis(user_id=user_id, cv_version_id=cv, status=Analysisstatus.COMPLETED)
        )
        first = client.post(
            f"/v1/analyses/{first_analysis_id}/generated-documents", headers=_headers(user_id)
        )
        assert first.status_code == 202
        _drain_generation_intake()

        second_analysis_id = asyncio.run(
            _seed_analysis(user_id=user_id, cv_version_id=cv, status=Analysisstatus.COMPLETED)
        )
        second = client.post(
            f"/v1/analyses/{second_analysis_id}/generated-documents", headers=_headers(user_id)
        )
        assert second.status_code == 429
        assert _drain_generation_intake() == []

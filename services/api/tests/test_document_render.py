import re
from io import BytesIO

from docx import Document
from docx.shared import Pt

from api.document_render import (
    render_markdown_to_docx,
    render_markdown_to_pdf,
    render_markdown_to_text,
)


def test_render_markdown_to_pdf_returns_pdf_bytes():
    pdf_bytes = render_markdown_to_pdf(markdown_content="Hello world")
    assert pdf_bytes.startswith(b"%PDF")


def test_render_markdown_to_docx_does_not_inject_title():
    docx_bytes = render_markdown_to_docx(markdown_content="Some paragraph text.")
    document = Document(BytesIO(docx_bytes))
    assert document.paragraphs[0].text == "Some paragraph text."
    assert document.paragraphs[0].style.name != "Title"


def test_render_markdown_to_pdf_handles_headings_and_bullets():
    markdown_content = "# Heading\n\nSome paragraph text.\n\n- First bullet\n- Second bullet\n"
    pdf_bytes = render_markdown_to_pdf(markdown_content=markdown_content)
    assert pdf_bytes.startswith(b"%PDF")
    assert len(pdf_bytes) > 0


def test_render_markdown_to_pdf_handles_empty_content():
    pdf_bytes = render_markdown_to_pdf(markdown_content="")
    assert pdf_bytes.startswith(b"%PDF")


def test_render_markdown_to_pdf_handles_characters_outside_latin1():
    pdf_bytes = render_markdown_to_pdf(
        markdown_content="Smart quotes: “quoted” and an em dash — here."
    )
    assert pdf_bytes.startswith(b"%PDF")


def test_render_markdown_to_docx_returns_zip_signature():
    docx_bytes = render_markdown_to_docx(markdown_content="Hello world")
    assert docx_bytes.startswith(b"PK")


def test_render_markdown_to_docx_handles_headings_and_bullets():
    markdown_content = "# Heading\n\n## Subheading\n\nSome paragraph text.\n\n- First bullet\n- Second bullet\n"
    docx_bytes = render_markdown_to_docx(markdown_content=markdown_content)
    assert docx_bytes.startswith(b"PK")
    assert len(docx_bytes) > 0


def test_render_markdown_to_docx_handles_empty_content():
    docx_bytes = render_markdown_to_docx(markdown_content="")
    assert docx_bytes.startswith(b"PK")


def test_render_markdown_to_text_strips_heading_syntax():
    markdown_content = "# Heading\n\n## Subheading\n\nSome paragraph text.\n"
    text = render_markdown_to_text(markdown_content=markdown_content)
    assert "# Heading" not in text
    assert "## Subheading" not in text
    assert "Heading" in text
    assert "Subheading" in text
    assert "Some paragraph text." in text


def test_render_markdown_to_text_leaves_bullets_as_is():
    markdown_content = "- First bullet\n* Second bullet\n"
    text = render_markdown_to_text(markdown_content=markdown_content)
    assert "- First bullet" in text
    assert "* Second bullet" in text


def test_render_markdown_to_text_handles_empty_content():
    assert render_markdown_to_text(markdown_content="") == ""


def test_render_markdown_to_text_strips_section_type_comment():
    markdown_content = "## Experience\n<!-- SectionType: EXPERIENCE -->\n\nDid things.\n"
    text = render_markdown_to_text(markdown_content=markdown_content)
    assert "SectionType" not in text
    assert "Experience" in text
    assert "Did things." in text


def test_render_markdown_to_docx_strips_section_type_comment():
    markdown_content = "## Experience\n<!-- SectionType: EXPERIENCE -->\n\nDid things.\n"
    docx_bytes = render_markdown_to_docx(markdown_content=markdown_content)
    document = Document(BytesIO(docx_bytes))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "SectionType" not in text
    assert "Experience" in text


_STYLE_PROFILE = {
    "layoutArchetype": "SINGLE_COLUMN",
    "fonts": {"name": "Georgia", "family": "serif"},
    "accentColor": "#112233",
    "margins": {"top": 40, "bottom": 40, "left": 40, "right": 40},
    "onePageFit": True,
    "photoAssetRef": None,
    "sections": {
        "EXPERIENCE": {
            "headingTreatment": {"font": {"name": "Impact", "family": "sans-serif"}, "color": "#FF0000"},
            "region": None,
        }
    },
}


def test_render_markdown_to_docx_applies_style_profile_heading_treatment():
    markdown_content = "## Experience\n<!-- SectionType: EXPERIENCE -->\n\nDid things.\n"
    docx_bytes = render_markdown_to_docx(
        markdown_content=markdown_content, style_profile=_STYLE_PROFILE
    )
    document = Document(BytesIO(docx_bytes))
    heading_paragraph = next(p for p in document.paragraphs if p.text == "Experience")
    run = heading_paragraph.runs[0]
    # "Impact" isn't in the curated cross-platform set, so it's substituted
    # by its family's ("sans-serif") entry rather than used literally (#100).
    assert run.font.name == "Arial"
    assert str(run.font.color.rgb) == "FF0000"

    body_paragraph = next(p for p in document.paragraphs if p.text == "Did things.")
    body_run = body_paragraph.runs[0]
    assert body_run.font.name == "Georgia"
    assert str(body_run.font.color.rgb) == "112233"


_SIDEBAR_MAIN_STYLE_PROFILE = {
    "layoutArchetype": "SIDEBAR_MAIN",
    "fonts": {"name": "Georgia", "family": "serif"},
    "accentColor": "#112233",
    "margins": {"top": 40, "bottom": 40, "left": 40, "right": 40},
    "onePageFit": False,
    "photoAssetRef": None,
    "sections": {
        "SKILLS": {
            "headingTreatment": {"font": {"name": "Georgia", "family": "serif"}, "color": "#112233"},
            "region": "SIDEBAR",
        },
        "EXPERIENCE": {
            "headingTreatment": {"font": {"name": "Impact", "family": "sans-serif"}, "color": "#FF0000"},
            "region": "MAIN",
        },
    },
}

_SIDEBAR_MAIN_MARKDOWN = (
    "## Skills\n<!-- SectionType: SKILLS -->\n\n- Python\n- SQL\n\n"
    "## Experience\n<!-- SectionType: EXPERIENCE -->\n\nDid things.\n"
)


def test_group_lines_by_region_splits_by_section_type():
    from api.document_render import _classify_lines, _group_lines_by_region

    lines = _classify_lines(_SIDEBAR_MAIN_MARKDOWN)
    sidebar_lines, main_lines = _group_lines_by_region(lines, _SIDEBAR_MAIN_STYLE_PROFILE)

    assert [line.text for line in sidebar_lines if line.text] == ["Skills", "Python", "SQL"]
    assert [line.text for line in main_lines if line.text] == ["Experience", "Did things."]


def test_group_lines_by_region_defaults_unassigned_sections_to_main():
    from api.document_render import _classify_lines, _group_lines_by_region

    markdown_content = "Some preamble.\n\n## Other\n<!-- SectionType: OTHER -->\n\nBody text.\n"
    lines = _classify_lines(markdown_content)
    sidebar_lines, main_lines = _group_lines_by_region(lines, _SIDEBAR_MAIN_STYLE_PROFILE)

    assert sidebar_lines == []
    assert [line.text for line in main_lines if line.text] == ["Some preamble.", "Other", "Body text."]


def test_render_markdown_to_docx_sidebar_main_splits_sections_into_columns():
    docx_bytes = render_markdown_to_docx(
        markdown_content=_SIDEBAR_MAIN_MARKDOWN, style_profile=_SIDEBAR_MAIN_STYLE_PROFILE
    )
    document = Document(BytesIO(docx_bytes))
    table = document.tables[0]
    sidebar_text = "\n".join(p.text for p in table.rows[0].cells[0].paragraphs)
    main_text = "\n".join(p.text for p in table.rows[0].cells[1].paragraphs)

    assert "Skills" in sidebar_text and "Python" in sidebar_text
    assert "Experience" not in sidebar_text
    assert "Experience" in main_text and "Did things." in main_text
    assert "Skills" not in main_text


def test_render_markdown_to_docx_sidebar_main_applies_heading_treatment_per_column():
    docx_bytes = render_markdown_to_docx(
        markdown_content=_SIDEBAR_MAIN_MARKDOWN, style_profile=_SIDEBAR_MAIN_STYLE_PROFILE
    )
    document = Document(BytesIO(docx_bytes))
    table = document.tables[0]
    main_heading = next(p for p in table.rows[0].cells[1].paragraphs if p.text == "Experience")
    run = main_heading.runs[0]
    # "Impact" isn't in the curated cross-platform set, so it's substituted
    # by its family's ("sans-serif") entry, same rule as SINGLE_COLUMN (#100).
    assert run.font.name == "Arial"
    assert str(run.font.color.rgb) == "FF0000"


def test_render_markdown_to_docx_single_column_unaffected_by_sidebar_main_support():
    markdown_content = "## Experience\n<!-- SectionType: EXPERIENCE -->\n\nDid things.\n"
    docx_bytes = render_markdown_to_docx(
        markdown_content=markdown_content, style_profile=_STYLE_PROFILE
    )
    document = Document(BytesIO(docx_bytes))
    assert document.tables == []
    heading_paragraph = next(p for p in document.paragraphs if p.text == "Experience")
    assert heading_paragraph.runs[0].font.name == "Arial"


def test_render_markdown_to_pdf_sidebar_main_renders_without_error():
    pdf_bytes = render_markdown_to_pdf(
        markdown_content=_SIDEBAR_MAIN_MARKDOWN, style_profile=_SIDEBAR_MAIN_STYLE_PROFILE
    )
    assert pdf_bytes.startswith(b"%PDF")


def test_render_markdown_to_pdf_applies_style_profile_without_error():
    markdown_content = "## Experience\n<!-- SectionType: EXPERIENCE -->\n\nDid things.\n"
    pdf_bytes = render_markdown_to_pdf(
        markdown_content=markdown_content, style_profile=_STYLE_PROFILE
    )
    assert pdf_bytes.startswith(b"%PDF")


def test_render_markdown_to_docx_substitutes_uncurated_body_font_by_family():
    style_profile = {**_STYLE_PROFILE, "fonts": {"name": "Papyrus", "family": "serif"}, "onePageFit": False}
    docx_bytes = render_markdown_to_docx(
        markdown_content="Some paragraph text.", style_profile=style_profile
    )
    document = Document(BytesIO(docx_bytes))
    body_paragraph = next(p for p in document.paragraphs if p.text == "Some paragraph text.")
    assert body_paragraph.runs[0].font.name == "Times New Roman"


def test_render_markdown_to_docx_keeps_curated_body_font_name():
    style_profile = {**_STYLE_PROFILE, "fonts": {"name": "Calibri", "family": "sans-serif"}, "onePageFit": False}
    docx_bytes = render_markdown_to_docx(
        markdown_content="Some paragraph text.", style_profile=style_profile
    )
    document = Document(BytesIO(docx_bytes))
    body_paragraph = next(p for p in document.paragraphs if p.text == "Some paragraph text.")
    assert body_paragraph.runs[0].font.name == "Calibri"


_LOREM_PARAGRAPH = (
    "Lorem ipsum dolor sit amet, consectetur adipiscing elit sed do eiusmod "
    "tempor incididunt ut labore et dolore magna aliqua."
)


def _pdf_page_count(pdf_bytes: bytes) -> int:
    match = re.search(rb"/Count (\d+)", pdf_bytes)
    assert match is not None
    return int(match.group(1))


def test_render_markdown_to_pdf_one_page_fit_reduces_to_stay_on_one_page():
    style_profile = {**_STYLE_PROFILE, "onePageFit": True}
    markdown_content = "\n\n".join([_LOREM_PARAGRAPH] * 28)
    pdf_bytes = render_markdown_to_pdf(
        markdown_content=markdown_content, style_profile=style_profile
    )
    assert _pdf_page_count(pdf_bytes) == 1


def test_render_markdown_to_pdf_one_page_fit_allows_overflow_when_too_long():
    style_profile = {**_STYLE_PROFILE, "onePageFit": True}
    markdown_content = "\n\n".join([_LOREM_PARAGRAPH] * 40)
    pdf_bytes = render_markdown_to_pdf(
        markdown_content=markdown_content, style_profile=style_profile
    )
    assert _pdf_page_count(pdf_bytes) > 1


def test_render_markdown_to_pdf_without_one_page_fit_does_not_reduce():
    style_profile = {**_STYLE_PROFILE, "onePageFit": False}
    markdown_content = "\n\n".join([_LOREM_PARAGRAPH] * 28)
    pdf_bytes = render_markdown_to_pdf(
        markdown_content=markdown_content, style_profile=style_profile
    )
    assert _pdf_page_count(pdf_bytes) == 2


def test_render_markdown_to_docx_one_page_fit_reduces_margins_and_font_for_long_content():
    style_profile = {**_STYLE_PROFILE, "onePageFit": True}
    markdown_content = "\n\n".join([_LOREM_PARAGRAPH] * 40)
    docx_bytes = render_markdown_to_docx(
        markdown_content=markdown_content, style_profile=style_profile
    )
    document = Document(BytesIO(docx_bytes))
    section = document.sections[0]
    assert section.top_margin == Pt(36)
    body_paragraph = next(p for p in document.paragraphs if p.text == _LOREM_PARAGRAPH)
    assert body_paragraph.runs[0].font.size < Pt(11)


def test_render_markdown_to_docx_one_page_fit_leaves_short_content_untouched():
    style_profile = {**_STYLE_PROFILE, "onePageFit": True}
    docx_bytes = render_markdown_to_docx(
        markdown_content="Short body text.", style_profile=style_profile
    )
    document = Document(BytesIO(docx_bytes))
    section = document.sections[0]
    assert section.top_margin != Pt(36)
    body_paragraph = next(p for p in document.paragraphs if p.text == "Short body text.")
    assert body_paragraph.runs[0].font.size is None


def test_render_markdown_to_docx_one_page_fit_budget_excludes_removed_title():
    # Exactly at the budget: with the removed title's characters still
    # counted (as before #105), this would have crossed it and triggered
    # reduction. It must not, now that the title no longer exists to count.
    style_profile = {**_STYLE_PROFILE, "onePageFit": True}
    markdown_content = "a" * 3500
    docx_bytes = render_markdown_to_docx(markdown_content=markdown_content, style_profile=style_profile)
    document = Document(BytesIO(docx_bytes))
    section = document.sections[0]
    assert section.top_margin != Pt(36)
